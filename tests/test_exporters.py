from openpyxl import Workbook, load_workbook

from cse_screening.exporters.workbook import (
    FINANCIAL_FORMAT,
    PER_SHARE_FORMAT,
    RATIO_FORMAT,
    SHEETS,
    _format_workbook,
    _neutral_markdown_summary,
    export_all,
)


def test_markdown_summary_includes_required_sections_and_neutral_extremes():
    rows = [
        {
            "Ticker": "AAA.N0000",
            "Data Confidence": "high",
            "Revenue": 200,
            "Net Profit": 20,
            "EPS": 2,
            "ROE": 0.2,
            "Operating Cash Flow": 25,
            "Free Cash Flow": 15,
            "DPS": 1,
            "Recent Volume": 1000,
            "P/E": 10,
            "P/B": 2,
            "Dividend Yield": 0.05,
            "Flags": "negative operating cash flow",
        },
        {
            "Ticker": "BBB.N0000",
            "Data Confidence": "review",
            "Revenue": 100,
            "Net Profit": None,
            "EPS": 1,
            "ROE": 0.1,
            "Operating Cash Flow": 5,
            "Free Cash Flow": 1,
            "DPS": 0,
            "Recent Volume": 10,
            "P/E": 20,
            "P/B": 1,
            "Dividend Yield": 0,
            "Flags": "",
        },
    ]
    review = [
        {
            "Ticker": "BBB.N0000",
            "Validation Rule": "missing_metric",
            "Reason for Uncertainty": "net profit unavailable",
        }
    ]
    text = _neutral_markdown_summary("2026Q3", "Capital Goods", rows, review)
    assert "## Sector medians" in text
    assert "## Highest and lowest observed values" in text
    assert "Revenue: highest AAA.N0000" in text
    assert "BBB.N0000:" in text and "Net Profit" in text
    assert "`missing_metric`" in text
    assert "no ranking" in text


def test_excel_financial_number_formats_are_metric_aware():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "01_Screening"
    sheet.append(["Current Price", "Revenue", "P/E", "ROE", "Flags"])
    sheet.append([12.5, -1000, 8.5, 0.12, "negative EPS"])
    _format_workbook(workbook)
    assert sheet["A2"].number_format == PER_SHARE_FORMAT
    assert sheet["B2"].number_format == FINANCIAL_FORMAT
    assert sheet["C2"].number_format == RATIO_FORMAT
    assert sheet["D2"].number_format == "0.00%;[Red]-0.00%"
    assert sheet.freeze_panes == "A2"
    assert sheet.auto_filter.ref == sheet.dimensions

    ratios = workbook.create_sheet("03_Ratios")
    ratios.append(["Metric", "Value"])
    ratios.append(["EBIT", 1000000])
    ratios.append(["P/E", 8.5])
    _format_workbook(workbook)
    assert ratios["B2"].number_format == FINANCIAL_FORMAT
    assert ratios["B3"].number_format == RATIO_FORMAT


def test_export_writes_specified_filenames_sheets_and_processed_tables(tmp_path):
    screening = [
        {
            "Ticker": "AAA.N0000",
            "Company": "AAA PLC",
            "P/E": 10,
            "P/B": 1,
            "ROE": 0.1,
            "Dividend Yield": 0.02,
            "Data Confidence": "high",
            "Income Period Basis": "TTM to 2026-06-30",
            "Flags": "",
        }
    ]
    raw = [{"Fact ID": "AAA:1", "Metric": "revenue", "Extracted Value": 5, "Unit": "LKR"}]
    outputs = export_all(
        "2026Q3",
        "Capital Goods",
        tmp_path / "reports",
        screening,
        raw,
        [{"Ticker": "AAA.N0000", "Metric": "P/E", "Value": 10}],
        [],
        [],
        [],
        [],
        [{"Ticker": "AAA.N0000", "Validation Rule": "missing_metric"}],
        [{"Ticker": "AAA.N0000"}],
        processed=tmp_path / "processed",
    )
    assert outputs["xlsx"].name == "Capital_Goods_Screening_2026_Q3.xlsx"
    assert outputs["csv"].name == "Capital_Goods_Screening_2026_Q3.csv"
    assert outputs["markdown"].name == "Capital_Goods_Summary_2026_Q3.md"
    assert outputs["review"].name == "manual_review.csv"
    assert outputs["processed_raw_facts"].exists()
    workbook = load_workbook(outputs["xlsx"])
    assert workbook.sheetnames == SHEETS
    assert set(SHEETS) >= {
        "01_Screening",
        "02_Raw_Data",
        "03_Ratios",
        "04_History",
        "05_Dividends",
        "06_Sources",
        "07_Flags",
        "08_Summary",
    }
    summary = outputs["markdown"].read_text(encoding="utf-8")
    assert "Sector median P/E: 10.00 (1 companies)" in summary
    assert "trailing-twelve-month basis: 1 companies" in summary
