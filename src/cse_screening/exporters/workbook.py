"""Required Excel, CSV, Markdown, and review exports."""

from __future__ import annotations

import json
import re
from pathlib import Path

import pandas as pd
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

SHEETS = [
    "00_Universe",
    "01_Screening",
    "02_Raw_Data",
    "03_Ratios",
    "04_History",
    "05_Dividends",
    "06_Sources",
    "07_Flags",
    "08_Summary",
]
PERCENT_COLUMNS = {
    "Revenue Growth YoY",
    "Net Profit Growth YoY",
    "EPS Growth YoY",
    "3Y Revenue CAGR",
    "3Y EPS CAGR",
    "ROE",
    "ROA",
    "Dividend Yield",
    "Payout Ratio",
    "Earnings Yield",
}
FINANCIAL_COLUMNS = {
    "Market Cap",
    "Revenue",
    "Operating Profit",
    "EBIT",
    "Net Profit",
    "Net Debt",
    "Retained Earnings",
    "Operating Cash Flow",
    "Free Cash Flow",
    "Numerator",
    "Denominator",
    "Current YTD",
    "Prior YTD",
}
PER_SHARE_COLUMNS = {
    "Current Price",
    "EPS",
    "DPS",
    "Book Value Per Share",
    "Voting DPS",
    "Non-Voting DPS",
}
INTEGER_COLUMNS = {
    "Ordinary Shares Outstanding",
    "Recent Volume",
    "Average Daily Volume",
    "Median Daily Volume",
    "Liquidity Trading Days",
}
RATIO_COLUMNS = {"P/E", "P/B", "Debt-to-Equity", "OCF / Net Profit"}
FINANCIAL_FORMAT = "#,##0;[Red](#,##0);-"
PER_SHARE_FORMAT = "#,##0.00;[Red](#,##0.00);-"
RATIO_FORMAT = '0.00"x";[Red](0.00"x");-'
INTEGER_FORMAT = "#,##0;[Red](#,##0);-"


def export_all(
    period: str,
    industry_group: str,
    output: Path,
    screening: list[dict],
    raw: list[dict],
    ratios: list[dict],
    history: list[dict],
    dividends: list[dict],
    sources: list[dict],
    flags: list[dict],
    review: list[dict],
    universe: list[dict],
) -> dict[str, Path]:
    output.mkdir(parents=True, exist_ok=True)
    safe_group = "_".join(re.findall(r"[A-Za-z0-9]+", industry_group))
    stem = f"{safe_group}_Screening_{period}"
    workbook_path = output / f"{stem}.xlsx"
    csv_path = output / f"{stem}.csv"
    summary_path = output / f"{safe_group}_Summary_{period}.md"
    universe_path = output / f"{safe_group}_Universe_{period}.csv"
    review_path = output / "manual_review.csv"
    warning_path = output / "extraction_warnings.jsonl"

    frames = {
        "00_Universe": pd.DataFrame(universe),
        "01_Screening": pd.DataFrame(screening),
        "02_Raw_Data": pd.DataFrame(raw),
        "03_Ratios": pd.DataFrame(ratios),
        "04_History": pd.DataFrame(history),
        "05_Dividends": pd.DataFrame(dividends),
        "06_Sources": pd.DataFrame(sources),
        "07_Flags": pd.DataFrame(flags),
        "08_Summary": _summary_frame(screening),
    }
    with pd.ExcelWriter(workbook_path, engine="openpyxl") as writer:
        for sheet in SHEETS:
            frames[sheet].to_excel(writer, sheet_name=sheet, index=False)
        _format_workbook(writer.book)
    frames["01_Screening"].to_csv(csv_path, index=False)
    frames["00_Universe"].to_csv(universe_path, index=False)
    pd.DataFrame(review).to_csv(review_path, index=False)
    warning_path.write_text(
        "".join(json.dumps(item, default=str) + "\n" for item in review), encoding="utf-8"
    )
    summary_path.write_text(
        _neutral_markdown_summary(period, industry_group, screening, review), encoding="utf-8"
    )
    return {
        "xlsx": workbook_path,
        "csv": csv_path,
        "markdown": summary_path,
        "universe_csv": universe_path,
        "review": review_path,
        "warnings": warning_path,
    }


def _format_workbook(workbook) -> None:
    for ws in workbook.worksheets:
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E78")
        for column in ws.columns:
            width = min(45, max(10, max(len(str(cell.value or "")) for cell in column) + 2))
            ws.column_dimensions[get_column_letter(column[0].column)].width = width
        headers = {cell.value: cell.column for cell in ws[1]}
        for name in PERCENT_COLUMNS:
            if name in headers:
                for cell in ws.iter_cols(min_col=headers[name], max_col=headers[name], min_row=2):
                    for item in cell:
                        item.number_format = "0.00%;[Red]-0.00%"
        for names, number_format in (
            (FINANCIAL_COLUMNS, FINANCIAL_FORMAT),
            (PER_SHARE_COLUMNS, PER_SHARE_FORMAT),
            (INTEGER_COLUMNS, INTEGER_FORMAT),
            (RATIO_COLUMNS, RATIO_FORMAT),
        ):
            for name in names:
                if name in headers:
                    for cells in ws.iter_cols(
                        min_col=headers[name], max_col=headers[name], min_row=2
                    ):
                        for item in cells:
                            item.number_format = number_format
        if ws.title == "02_Raw_Data":
            _format_value_by_unit(ws, headers, "Extracted Value", "Unit")
            _format_value_by_unit(ws, headers, "Original Value", "Original Unit")
        elif ws.title == "04_History":
            _format_value_by_unit(ws, headers, "Value", "Unit")
        elif ws.title == "03_Ratios" and "Value" in headers and "Metric" in headers:
            for row_number in range(2, ws.max_row + 1):
                metric = ws.cell(row_number, headers["Metric"]).value
                if metric in PERCENT_COLUMNS:
                    number_format = "0.00%;[Red]-0.00%"
                elif metric in RATIO_COLUMNS:
                    number_format = RATIO_FORMAT
                else:
                    number_format = FINANCIAL_FORMAT
                ws.cell(row_number, headers["Value"]).number_format = number_format
        elif ws.title == "08_Summary" and "Value" in headers and "Metric" in headers:
            for row_number in range(2, ws.max_row + 1):
                metric = str(ws.cell(row_number, headers["Metric"]).value or "")
                value_cell = ws.cell(row_number, headers["Value"])
                if metric.endswith(("ROE", "Dividend Yield")):
                    value_cell.number_format = "0.00%;[Red]-0.00%"
                elif metric.endswith(("P/E", "P/B")):
                    value_cell.number_format = RATIO_FORMAT
                else:
                    value_cell.number_format = INTEGER_FORMAT
        if ws.title == "01_Screening" and ws.max_row > 1:
            for name in ("P/E", "P/B", "ROE", "Debt-to-Equity", "Dividend Yield"):
                if name in headers:
                    letter = get_column_letter(headers[name])
                    ws.conditional_formatting.add(
                        f"{letter}2:{letter}{ws.max_row}",
                        ColorScaleRule(
                            start_type="min",
                            start_color="F8696B",
                            mid_type="percentile",
                            mid_value=50,
                            mid_color="FFEB84",
                            end_type="max",
                            end_color="63BE7B",
                        ),
                    )
            negative_fill = PatternFill("solid", fgColor="FFC7CE")
            for name in ("EPS", "Operating Cash Flow", "Free Cash Flow"):
                if name in headers:
                    letter = get_column_letter(headers[name])
                    ws.conditional_formatting.add(
                        f"{letter}2:{letter}{ws.max_row}",
                        CellIsRule(operator="lessThan", formula=["0"], fill=negative_fill),
                    )
            if "Flags" in headers:
                letter = get_column_letter(headers["Flags"])
                ws.conditional_formatting.add(
                    f"{letter}2:{letter}{ws.max_row}",
                    CellIsRule(
                        operator="notEqual",
                        formula=['""'],
                        fill=PatternFill("solid", fgColor="FFF2CC"),
                    ),
                )


def _format_value_by_unit(ws, headers: dict, value_name: str, unit_name: str) -> None:
    if value_name not in headers or unit_name not in headers:
        return
    for row_number in range(2, ws.max_row + 1):
        unit = str(ws.cell(row_number, headers[unit_name]).value or "").casefold()
        cell = ws.cell(row_number, headers[value_name])
        if "/share" in unit or "dps" in unit:
            cell.number_format = PER_SHARE_FORMAT
        elif "share" in unit and "lkr" not in unit:
            cell.number_format = INTEGER_FORMAT
        else:
            cell.number_format = FINANCIAL_FORMAT


def _summary_frame(rows: list[dict]) -> pd.DataFrame:
    metrics = ["P/E", "P/B", "ROE", "Dividend Yield"]
    data = []
    frame = pd.DataFrame(rows)
    for metric in metrics:
        series = (
            pd.to_numeric(frame[metric], errors="coerce")
            if metric in frame
            else pd.Series(dtype=float)
        )
        data.append(
            {
                "Metric": f"Sector median {metric}",
                "Value": series.median() if not series.empty else None,
            }
        )
    data.insert(
        0,
        {
            "Metric": "Companies with high data confidence",
            "Value": sum(row.get("Data Confidence") == "high" for row in rows),
        },
    )
    data.insert(0, {"Metric": "Companies processed", "Value": len(rows)})
    return pd.DataFrame(data)


def _markdown_summary(period: str, rows: list[dict], review: list[dict]) -> str:
    frame = pd.DataFrame(rows)
    lines = [
        f"# Capital Goods screening summary — {period}",
        "",
        f"Companies processed: **{len(rows)}**.",
        "",
    ]
    for metric in ("P/E", "P/B", "ROE", "Dividend Yield"):
        series = (
            pd.to_numeric(frame[metric], errors="coerce").dropna()
            if metric in frame
            else pd.Series(dtype=float)
        )
        value = series.median() if not series.empty else None
        rendered = "NA" if value is None else f"{value:.4f}"
        lines.append(f"- Sector median {metric}: {rendered}")
    missing = [row["Ticker"] for row in rows if row.get("Data Confidence") != "high"]
    lines += [
        "",
        "## Missing data and warnings",
        "",
        f"Companies requiring review: {', '.join(missing) or 'None'}.",
        f"Manual-review items: {len(review)}.",
        "",
        "This report contains no ranking, score, selection, or investment recommendation.",
        "",
    ]
    return "\n".join(lines)


def _neutral_markdown_summary(
    period: str, industry_group: str, rows: list[dict], review: list[dict]
) -> str:
    frame = pd.DataFrame(rows)
    high_confidence = sum(row.get("Data Confidence") == "high" for row in rows)
    lines = [
        f"# {industry_group} screening summary - {period}",
        "",
        f"Companies processed: **{len(rows)}**.",
        f"Companies with high data confidence: **{high_confidence}**.",
        f"Companies requiring review: **{len(rows) - high_confidence}**.",
        "",
        "## Sector medians",
        "",
    ]
    for metric in ("P/E", "P/B", "ROE", "Dividend Yield"):
        series = (
            pd.to_numeric(frame[metric], errors="coerce").dropna()
            if metric in frame
            else pd.Series(dtype=float)
        )
        value = series.median() if not series.empty else None
        rendered = _render_metric(metric, value)
        lines.append(f"- Sector median {metric}: {rendered}")

    observed_metrics = (
        "Revenue",
        "Net Profit",
        "EPS",
        "ROE",
        "Operating Cash Flow",
        "Free Cash Flow",
        "DPS",
        "Recent Volume",
        "Average Daily Volume",
    )
    lines += ["", "## Highest and lowest observed values", ""]
    for metric in observed_metrics:
        if metric not in frame:
            continue
        values = pd.to_numeric(frame[metric], errors="coerce").dropna()
        if values.empty:
            lines.append(f"- {metric}: NA")
            continue
        high_index, low_index = values.idxmax(), values.idxmin()
        lines.append(
            f"- {metric}: highest {frame.loc[high_index, 'Ticker']} "
            f"({_render_metric(metric, values.loc[high_index])}); lowest "
            f"{frame.loc[low_index, 'Ticker']} ({_render_metric(metric, values.loc[low_index])})"
        )

    completeness_metrics = (
        "Current Price",
        "Market Cap",
        "Revenue",
        "Net Profit",
        "EPS",
        "ROE",
        "ROA",
        "Debt-to-Equity",
        "Operating Cash Flow",
        "Free Cash Flow",
        "P/E",
        "P/B",
        "DPS",
        "Dividend Yield",
    )
    missing_rows = []
    for row in rows:
        missing = [metric for metric in completeness_metrics if row.get(metric) is None]
        if missing:
            missing_rows.append(f"- {row['Ticker']}: {', '.join(missing)}")
    lines += ["", "## Missing data", ""]
    lines.extend(missing_rows or ["- None"])

    flag_counts: dict[str, int] = {}
    for row in rows:
        for flag in (item for item in str(row.get("Flags") or "").split("; ") if item):
            flag_counts[flag] = flag_counts.get(flag, 0) + 1
    lines += ["", "## Anomaly flag counts", ""]
    lines.extend(
        [f"- {flag}: {count}" for flag, count in sorted(flag_counts.items())]
        or ["- No anomalies flagged"]
    )

    warning_groups: dict[tuple[str, str], set[str]] = {}
    for item in review:
        rule = str(item.get("Validation Rule") or "unspecified")
        reason = str(item.get("Reason for Uncertainty") or "No reason supplied")
        warning_groups.setdefault((rule, reason), set()).add(str(item.get("Ticker") or "unknown"))
    lines += [
        "",
        "## Extraction warnings",
        "",
        f"Manual-review items: {len(review)}.",
    ]
    for (rule, reason), tickers in sorted(warning_groups.items()):
        lines.append(
            f"- `{rule}` ({len(tickers)} companies: {', '.join(sorted(tickers))}): {reason}"
        )
    if not warning_groups:
        lines.append("- None")
    lines += [
        "",
        (
            "Values above are descriptive observations only. This report contains no ranking, "
            "score, selection, or investment recommendation."
        ),
        "",
    ]
    return "\n".join(lines)


def _render_metric(metric: str, value) -> str:
    if value is None or pd.isna(value):
        return "NA"
    if metric in PERCENT_COLUMNS:
        return f"{float(value):.2%}"
    if metric in PER_SHARE_COLUMNS or metric in RATIO_COLUMNS:
        return f"{float(value):,.2f}"
    return f"{float(value):,.0f}"
